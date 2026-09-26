"""
Genera un CV de BORRADOR en inglés, hardware-first, formato EE. UU. y compatible con ATS, con los datos del sitio.
Desde el 25 sep 2026 el sitio ya NO lo publica: sirve las 4 hojas de vida propias de public/assets/cv/.
Queda como herramienta local (p. ej. para comparar datos o armar una versión nueva).

Salidas:
  docs/cv/Nicolas_Gomez_CV.docx          versión limpia (sin marcas de pendiente)
  docs/cv/Nicolas_Gomez_CV_WORKING.docx  versión de trabajo con ⚠ resaltado en amarillo
  docs/cv/Nicolas_Gomez_CV.pdf           PDF limpio (exportado con Microsoft Word)

El dominio sale de site.config.js; correo, teléfono, LinkedIn y promedio salen de PROFILE en src/data/content.js.
Uso:  python scripts/build_cv.py
"""

import re
import subprocess
import sys
from pathlib import Path
from urllib.parse import unquote

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
CONTENT = (ROOT / 'src/data/content.js').read_text(encoding='utf-8')
SITE_URL = re.search(r"SITE_URL\s*=\s*'([^']+)'", (ROOT / 'site.config.js').read_text(encoding='utf-8')).group(1)


def profile(key):
    """Lee un campo de PROFILE en content.js (única fuente de verdad)."""
    return re.search(rf"^  {key}:\s*'([^']+)'", CONTENT, re.M).group(1)


EMAIL = profile('email')
PHONE = profile('phone')
GPA = f"{profile('gpa')} / {profile('gpaScale')}"
GITHUB = 'github.com/NicoGomez4262'
LINKEDIN_URL = profile('linkedin')
LINKEDIN = unquote(LINKEDIN_URL).replace('https://www.', '')

INK = RGBColor(0x11, 0x18, 0x22)
DIM = RGBColor(0x44, 0x4E, 0x5A)
ACCENT = RGBColor(0x0E, 0x5A, 0x70)

PAGE_W, MARGIN_X = 8.5, 0.6
TEXT_W = PAGE_W - 2 * MARGIN_X


# --------------------------------------------------------------------------- helpers
def set_border_bottom(paragraph, color='0E5A70', size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for k, v in {'w:val': 'single', 'w:sz': str(size), 'w:space': '1', 'w:color': color}.items():
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def add_hyperlink(paragraph, url, text, size=None, color=ACCENT):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), r_id)
    run = OxmlElement('w:r')
    r_pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color')
    c.set(qn('w:val'), str(color))
    r_pr.append(c)
    if size:
        sz = OxmlElement('w:sz')
        sz.set(qn('w:val'), str(int(size * 2)))
        r_pr.append(sz)
    run.append(r_pr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def spacing(p, before=0, after=0, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


class CV:
    def __init__(self, working):
        self.working = working
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(MARGIN_X)
        sec.top_margin, sec.bottom_margin = Inches(0.42), Inches(0.38)

        normal = self.doc.styles['Normal']
        normal.font.name = 'Calibri'
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        normal.font.size = Pt(9.5)
        normal.font.color.rgb = INK

        lb = self.doc.styles['List Bullet']
        lb.font.name = 'Calibri'
        lb.font.size = Pt(9.5)

    # Marca de pendiente: solo en la versión de trabajo.
    def pending(self, p, text):
        if not self.working:
            return
        r = run(p, f' ⚠ PENDING: {text}', bold=True, size=9)
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def pending_line(self, text):
        if not self.working:
            return
        p = self.doc.add_paragraph()
        spacing(p, 0, 1)
        self.pending(p, text)

    def heading(self, text):
        p = self.doc.add_paragraph()
        spacing(p, 6, 2)
        run(p, text.upper(), bold=True, size=10.5, color=ACCENT)
        set_border_bottom(p)
        return p

    def entry(self, left_bold, left_rest, right, italic_line=None):
        p = self.doc.add_paragraph()
        spacing(p, 3, 0)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
        run(p, left_bold, bold=True)
        if left_rest:
            run(p, left_rest, color=DIM)
        run(p, '\t' + right, color=DIM)
        if italic_line:
            q = self.doc.add_paragraph()
            spacing(q, 0, 1)
            run(q, italic_line, italic=True, color=DIM)
        return p

    def bullet(self, text, bold_prefix=None):
        p = self.doc.add_paragraph(style='List Bullet')
        spacing(p, 0, 1, 1.0)
        pf = p.paragraph_format
        pf.left_indent = Inches(0.2)
        pf.first_line_indent = Inches(-0.14)
        if bold_prefix:
            run(p, bold_prefix, bold=True)
        run(p, text)
        return p

    def skill(self, label, value):
        p = self.doc.add_paragraph()
        spacing(p, 0, 1)
        pf = p.paragraph_format
        pf.left_indent = Inches(0)
        run(p, f'{label}: ', bold=True)
        run(p, value)
        return p

    # ----------------------------------------------------------------------- content
    def build(self):
        d = self.doc

        # Header
        p = d.add_paragraph()
        spacing(p, 0, 0)
        run(p, 'NICOLÁS ANDRÉS GÓMEZ RAMÍREZ', bold=True, size=18, color=INK)

        p = d.add_paragraph()
        spacing(p, 1, 2)
        run(p, 'Electronic Engineering Student  |  Embedded Systems & PCB Design', size=10.5, color=ACCENT, bold=True)

        p = d.add_paragraph()
        spacing(p, 0, 0)
        run(p, f'Bogotá, Colombia | {PHONE} (WhatsApp) | ', color=DIM)
        add_hyperlink(p, f'mailto:{EMAIL}', EMAIL)
        p = d.add_paragraph()
        spacing(p, 0, 0)
        add_hyperlink(p, SITE_URL, SITE_URL.replace('https://', ''))
        run(p, ' | ', color=DIM)
        add_hyperlink(p, f'https://{GITHUB}', GITHUB)
        run(p, ' | ', color=DIM)
        add_hyperlink(p, LINKEDIN_URL, LINKEDIN)

        # Education
        self.heading('Education')
        self.entry('Pontificia Universidad Javeriana', ', Bogotá, Colombia', 'Jan 2023 – Expected 2027')
        p = d.add_paragraph()
        spacing(p, 0, 1)
        run(p, 'B.S. in Electronic Engineering', italic=True)
        run(p, f'  |  GPA: {GPA}')
        self.pending_line('Relevant coursework (exact course names), e.g. Digital Electronics, Signals & Systems, Processor-Based Systems Design')

        # Technical skills — hardware first
        self.heading('Technical Skills')
        self.skill('PCB & Hardware Design', 'Altium Designer (schematic capture, PCB layout; 4+ custom PCBs to date), LTspice, 3D design & printing')
        self.skill('Embedded Systems', 'C, C++, MicroPython; PIC (MPLAB X, MCC), Arduino Uno / Mega 2560, Raspberry Pi (Python, GPIO); '
                   'SPI, UART, I²C, 1-Wire, PWM, MQTT; finite-state machines (FSM), interrupts, timing analysis')
        self.skill('Sensors & Actuation', 'IR, ADXL345 accelerometer, 12-bit SPI ADC/DAC; SSD1306 OLED, I²C LCD, LED matrices; '
                   'DC motor control (H-bridge, PWM), servo control')
        self.skill('Digital Design', 'VHDL, Verilog; Intel Quartus, ModelSim testbenches; Cyclone IV E FPGA (DE2-115), VGA timing')
        self.skill('Signals', 'MATLAB; FIR/IIR filter design, FIR on MCU; analog and anti-aliasing filter design; correlation-based detection')
        self.skill('Software & Tools', 'Python, JavaScript/TypeScript, React, Next.js, Firebase, PostgreSQL; Git, Linux, Windows, Vercel')
        self.skill('AI', 'Claude Code, Codex (coding agents); Higgsfield (generative media)')
        b = self.skill('Lab & Instrumentation', 'Oscilloscope, digital multimeter, bench power supply, function generator, soldering')
        self.pending(b, 'THT, SMD or both?')

        # Experience: academic first, then professional
        self.heading('Academic Experience')
        self.entry('Teaching Assistant — Processor-Based Systems Design', ', Pontificia Universidad Javeriana', 'Feb 2026 – Present')
        b = self.bullet('Designed and developed custom PCBs in Altium Designer, from schematic capture to layout, for hands-on '
                        'experimentation with course concepts.')
        self.pending(b, 'board function, MCU, layers, units built; students per semester (only if exact)')
        self.bullet('Resolve office-hours questions on embedded C, finite-state machines, interrupts and timing analysis; '
                    'support embedded-systems and microprocessor course projects.')
        self.entry('Teaching Assistant — Arduino 101', ', Pontificia Universidad Javeriana', 'Jan 2025 – Jun 2025')
        self.bullet('Guided students through sensors, actuators and peripherals on Arduino; mentored projects from design to '
                    'implementation and debugged embedded C (memory management, optimization).')

        self.heading('Professional Experience')
        self.entry('Academic Tutor — Programa Tú', ', Bogotá, Colombia', 'Sep 2024 – Present')
        self.bullet('Provide personalized tutoring in physics, mathematics, critical thinking and language to students at '
                    'Instituto Alberto Merani (22+ contracted hours).')
        self.entry('Freelance Web Developer', ', Independent businesses', 'Jul 2026 – Present')
        self.bullet('Develop web applications for independent businesses, from requirements definition through implementation '
                    'and delivery of solutions tailored to their needs.')

        # Hardware projects, in the site's order
        self.heading('Hardware Projects')
        self.entry('PROTEO — Raspberry Pi IoT Node for a Smart Luggage Locker', '  |  Python, MQTT, Firebase', '2025')
        b = self.bullet('IR sensor published over MQTT (QoS 1) on change, servo latch driven with microsecond pulses (pigpio) and a retained '
                        'Last-Will status topic; the final prototype syncs a solenoid lock, door state and camera photos with a web '
                        'dashboard through Firestore.')
        self.pending(b, 'course, team and your role')

        self.entry('Real-Time FIR Filtering on a PIC Microcontroller', '  |  C, SPI, UART', '2025')
        b = self.bullet('Two-channel chain: MAX11666 12-bit ADC → order-12/14 FIR filters in C → MCP4822 12-bit DAC on one shared SPI '
                        'bus, with a UART command parser and FSM to switch filter, input and routing at runtime.')
        self.pending(b, 'sample rate, exact PIC part number, solo or team')

        self.entry('VOLLEY-PONG — Two-Player Game on an FPGA (DE2-115)', '  |  VHDL, team of 3', '2024')
        self.bullet('800×600 VGA sync generator verified in ModelSim, fixed-point ball physics with gravity and collisions, boost '
                    'meter and scoring in pure VHDL; joystick inputs through LM393 comparators.')

        self.entry('CORRELACIUM-LEVIOSA — Gesture Recognition by Correlation', '  |  MicroPython, I²C, team of 2', '2024')
        self.bullet('ADXL345 sampled at 100 Hz into 100-sample windows; four spell gestures matched by correlation against recorded '
                    'templates with per-class thresholds; result on an SSD1306 OLED.')

        self.entry('DC Motor Control for a Conveyor Belt', '  |  Arduino, C', '2025')
        self.bullet('Direction and speed from one potentiometer via an L298N H-bridge (PWM), with a dead zone and a friction-aware minimum duty.')

        self.entry('DreamSnake — Two-Player Game on Arduino Mega 2560', '  |  C++, team of 3', '2023')
        self.bullet('Two MAX7219 8×8 LED matrices and a 20×4 I²C LCD, two-player buttons and an FSM-sequenced tune.')

        # Languages
        self.heading('Languages')
        p = d.add_paragraph()
        spacing(p, 0, 0)
        run(p, 'Spanish: ', bold=True)
        run(p, 'Native  |  ')
        run(p, 'English: ', bold=True)
        run(p, 'Upper-Intermediate (B2), Berlitz English Program, 2026')
        return self


def to_pdf(docx_path, pdf_path):
    """Exporta con Microsoft Word (texto real, legible por ATS)."""
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx_path}'); "
        f"$d.ExportAsFixedFormat('{pdf_path}', 17); "
        "$d.Close($false); $w.Quit()"
    )
    subprocess.run(['powershell', '-NoProfile', '-Command', ps], check=True)


if __name__ == '__main__':
    out = ROOT / 'docs' / 'cv'
    out.mkdir(parents=True, exist_ok=True)
    clean = out / 'Nicolas_Gomez_CV.docx'
    working = out / 'Nicolas_Gomez_CV_WORKING.docx'
    CV(working=False).build().doc.save(clean)
    CV(working=True).build().doc.save(working)
    pdf = out / 'Nicolas_Gomez_CV.pdf'
    if '--no-pdf' not in sys.argv:
        to_pdf(str(clean), str(pdf))
    print('SITE_URL', SITE_URL, '| EMAIL', EMAIL, '| GPA', GPA, '| LINKEDIN', LINKEDIN)
    print('wrote', clean, working, pdf)
